"""RAG-Source doit traiter n'importe quel document, pas un corpus particulier.

Ces tests prennent volontairement le contre-pied du corpus de développement :
notice d'électroménager en anglais, mise en page sur deux colonnes, page scannée,
export CSV, page web enregistrée, document Word.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from rag_source.config import OcrMode
from rag_source.domain import SectionKind
from rag_source.ingest.loaders.base import LoaderError
from rag_source.ingest.loaders.delimited import DelimitedLoader
from rag_source.ingest.loaders.docx import DocxLoader
from rag_source.ingest.loaders.html import HtmlLoader
from rag_source.ingest.loaders.pdf import PdfLoader
from rag_source.ingest.loaders.plaintext import PlainTextLoader


class TestPlainText:
    def test_underlined_and_numbered_headings(self, tmp_path: Path) -> None:
        content = (
            "MICROWAVE OVEN USER MANUAL\n"
            "==========================\n\n"
            "1. Safety instructions\n\n"
            "Do not operate the oven when empty.\n\n"
            "1.1 Grounding\n\n"
            "This appliance must be grounded.\n"
        )
        path = tmp_path / "manual.txt"
        path.write_text(content, encoding="utf-8")

        extracted = PlainTextLoader().load(path)
        assert extracted.title == "MICROWAVE OVEN USER MANUAL"
        paths = [s.heading_path for s in extracted.sections]
        assert ("MICROWAVE OVEN USER MANUAL", "1. Safety instructions") in paths
        assert ("MICROWAVE OVEN USER MANUAL", "1. Safety instructions", "1.1 Grounding") in paths

    def test_non_utf8_encoding(self, tmp_path: Path) -> None:
        path = tmp_path / "notes.txt"
        path.write_bytes("Réglage du minuteur à 3 minutes.\n".encode("cp1252"))
        assert "minuteur" in PlainTextLoader().load(path).sections[0].text


class TestDelimited:
    def test_semicolon_export_becomes_records(self, tmp_path: Path) -> None:
        path = tmp_path / "pieces.csv"
        path.write_text(
            "Part number;Description;Price\nMW-104;Turntable motor;24.90\nMW-207;Door latch;8.50\n",
            encoding="utf-8",
        )
        extracted = DelimitedLoader().load(path)
        records = [s for s in extracted.sections if s.kind is SectionKind.RECORD]
        assert len(records) == 2
        assert records[0].metadata["rec_id"] == "MW-104"  # colonne reconnue sans libellé français
        assert "Description : Turntable motor" in records[0].text

    def test_tsv(self, tmp_path: Path) -> None:
        path = tmp_path / "codes.tsv"
        path.write_text("Code\tMeaning\nE01\tDoor open\nE02\tSensor failure\n", encoding="utf-8")
        assert len(DelimitedLoader().load(path).sections) == 2

    def test_empty_file(self, tmp_path: Path) -> None:
        path = tmp_path / "vide.csv"
        path.write_text("", encoding="utf-8")
        with pytest.raises(LoaderError, match="vide"):
            DelimitedLoader().load(path)


class TestHtml:
    PAGE = """<html><head><title>Microwave FAQ</title></head><body>
      <nav>Menu principal</nav>
      <script>trackVisitor();</script>
      <h1>Troubleshooting</h1>
      <p>The oven does not heat.</p>
      <h2>Error codes</h2>
      <ul><li>E01: door open</li><li>E02: sensor failure</li></ul>
      <table><tr><th>Code</th><th>Action</th></tr>
      <tr><td>E01</td><td>Close the door</td></tr></table>
      <footer>Copyright 2026</footer>
    </body></html>"""

    def test_structure_and_noise_removal(self, tmp_path: Path) -> None:
        path = tmp_path / "faq.html"
        path.write_text(self.PAGE, encoding="utf-8")

        extracted = HtmlLoader().load(path)
        assert extracted.title == "Microwave FAQ"
        text = "\n".join(s.text for s in extracted.sections)
        assert "Menu principal" not in text and "trackVisitor" not in text
        assert "Copyright" not in text
        assert "- E01: door open" in text

        tables = [s for s in extracted.sections if s.kind is SectionKind.TABLE]
        assert len(tables) == 1
        assert tables[0].heading_path == ("Troubleshooting", "Error codes")
        assert "Close the door" in tables[0].text


class TestDocx:
    def test_headings_lists_and_tables(self, tmp_path: Path) -> None:
        import docx

        document = docx.Document()
        document.add_heading("Installation", level=1)
        document.add_paragraph("Place the oven on a flat surface.")
        document.add_heading("Clearances", level=2)
        document.add_paragraph("20 cm above the appliance", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Side"
        table.cell(0, 1).text = "Minimum"
        table.cell(1, 0).text = "Top"
        table.cell(1, 1).text = "20 cm"
        path = tmp_path / "manual.docx"
        document.save(str(path))

        extracted = DocxLoader().load(path)
        paths = [s.heading_path for s in extracted.sections]
        assert ("Installation",) in paths
        assert ("Installation", "Clearances") in paths
        assert any("- 20 cm above the appliance" in s.text for s in extracted.sections)
        tables = [s for s in extracted.sections if s.kind is SectionKind.TABLE]
        assert len(tables) == 1 and "Minimum" in tables[0].text


class TestPdfLayouts:
    def test_two_column_reading_order(self, tmp_path: Path) -> None:
        """Une notice sur deux colonnes doit se lire colonne par colonne.

        Un tri par position verticale seule entrelacerait les deux colonnes et
        produirait un texte incohérent.
        """
        import pymupdf

        document = pymupdf.open()
        page = document.new_page()
        left = ["Press START to begin.", "The turntable rotates.", "A beep ends the cycle."]
        right = ["Do not heat sealed jars.", "Keep the vents clear.", "Clean with a damp cloth."]
        y = 120.0
        for line_left, line_right in zip(left, right, strict=True):
            page.insert_text((60, y), line_left, fontsize=11, fontname="Times-Roman")
            page.insert_text((320, y), line_right, fontsize=11, fontname="Times-Roman")
            y += 40
        path = tmp_path / "two-column.pdf"
        document.save(path)
        document.close()

        text = "\n".join(
            s.text for s in PdfLoader(ocr_mode=OcrMode.OFF, ocr_languages="eng").load(path).sections
        )
        assert text.index("A beep ends the cycle") < text.index("Do not heat sealed jars")

    def test_full_width_heading_separates_columns(self, tmp_path: Path) -> None:
        """Un titre pleine largeur reste à sa place entre deux blocs de colonnes."""
        import pymupdf

        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((60, 80), "CLEANING AND MAINTENANCE", fontsize=16, fontname="Times-Bold")
        y = 130.0
        for index in range(3):
            page.insert_text((60, y), f"Left line {index}.", fontsize=11, fontname="Times-Roman")
            page.insert_text((320, y), f"Right line {index}.", fontsize=11, fontname="Times-Roman")
            y += 40
        path = tmp_path / "heading.pdf"
        document.save(path)
        document.close()

        sections = PdfLoader(ocr_mode=OcrMode.OFF, ocr_languages="eng").load(path).sections
        assert all(s.heading_path == ("CLEANING AND MAINTENANCE",) for s in sections)


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="Tesseract non installé")
class TestOcr:
    def test_scanned_page_is_recognized(self, tmp_path: Path) -> None:
        """Une page sans couche texte (scan) doit être lue par OCR."""
        import pymupdf

        # Page rendue en image : plus aucun texte sélectionnable.
        source = pymupdf.open()
        page = source.new_page()
        page.insert_text((72, 200), "DEFROST SETTING", fontsize=28, fontname="Helvetica-Bold")
        pixmap = page.get_pixmap(dpi=200)
        width, height = page.rect.width, page.rect.height
        source.close()

        scanned = pymupdf.open()
        image_page = scanned.new_page(width=width, height=height)
        image_page.insert_image(image_page.rect, pixmap=pixmap)
        path = tmp_path / "scan.pdf"
        scanned.save(path)
        scanned.close()

        with pytest.raises(LoaderError, match="scanné"):
            PdfLoader(ocr_mode=OcrMode.OFF, ocr_languages="eng").load(path)

        extracted = PdfLoader(ocr_mode=OcrMode.AUTO, ocr_languages="eng").load(path)
        text = " ".join(s.text for s in extracted.sections).upper()
        assert "DEFROST" in text
        assert any("OCR" in warning for warning in extracted.warnings)
